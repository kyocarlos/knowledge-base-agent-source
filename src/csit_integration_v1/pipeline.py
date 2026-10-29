"""Typed extraction -> validation -> deterministic writers. No Agent storage tools."""
import csv,hashlib,json,math,zipfile,io,uuid,re
from datetime import datetime,timezone
from pathlib import Path
from typing import Any
from pydantic import BaseModel,ConfigDict,Field,field_validator
from .errors import IntegrationError

PARSER_VERSION="km-csit-parser-v1"
EXTRACTION_VERSION="km-csit-schema-v1"
FORMAT_MIME={
 ".txt":"text/plain",".csv":"text/csv",".pdf":"application/pdf",
 ".docx":"application/vnd.openxmlformats-officedocument.wordprocessingml.document",
 ".xlsx":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
def stable_id(*parts):
    return str(uuid.uuid5(uuid.NAMESPACE_URL,json.dumps(parts,ensure_ascii=False,separators=(",",":"))))
def lineage(job):
    p=job["request_payload"]
    return {k:p[k] for k in ("sourceSystem","project","case","task","documentId","version")}
def validate_file(path,job,settings):
    suffix=Path(job["file_name"]).suffix.lower()
    if suffix not in FORMAT_MIME:raise IntegrationError("UNSUPPORTED_DOCUMENT_FORMAT","Document format unsupported")
    if job["content_type"]!=FORMAT_MIME[suffix]:raise IntegrationError("UNSUPPORTED_DOCUMENT_FORMAT","Extension and MIME type incompatible")
    if path.stat().st_size==0:raise IntegrationError("CORRUPT_DOCUMENT","Empty source document")
    with path.open("rb") as f:head=f.read(16)
    if head.startswith((b"MZ",b"\x7fELF")):raise IntegrationError("INVALID_FILE_SIGNATURE","Executable content rejected")
    if suffix==".pdf" and not head.startswith(b"%PDF-"):raise IntegrationError("INVALID_FILE_SIGNATURE","PDF signature invalid")
    if suffix in (".docx",".xlsx"):
        try:
            with zipfile.ZipFile(path) as z:
                infos=z.infolist()
                if len(infos)>settings.zip_max_members or sum(x.file_size for x in infos)>settings.zip_max_expanded:
                    raise IntegrationError("ARCHIVE_LIMIT_EXCEEDED","Archive expansion limits exceeded")
                for x in infos:
                    if x.file_size/max(1,x.compress_size)>settings.zip_max_ratio:
                        raise IntegrationError("ARCHIVE_LIMIT_EXCEEDED","Archive compression ratio exceeded")
                    if x.filename.startswith(("/","\\")) or ".." in x.filename.replace("\\","/").split("/"):
                        raise IntegrationError("UNSAFE_ARCHIVE","Archive member path unsafe")
                names=set(z.namelist())
                required="word/document.xml" if suffix==".docx" else "xl/workbook.xml"
                if required not in names or "[Content_Types].xml" not in names:
                    raise IntegrationError("CORRUPT_DOCUMENT","Office document structure invalid")
        except (zipfile.BadZipFile,RuntimeError):
            raise IntegrationError("CORRUPT_DOCUMENT","Office archive invalid") from None
    return suffix

class Entity(BaseModel):
    model_config=ConfigDict(extra="forbid")
    key:str=Field(min_length=1,max_length=256)
    kind:str=Field(min_length=1,max_length=64)
    name:str=Field(min_length=1,max_length=1024)
class Relation(BaseModel):
    model_config=ConfigDict(extra="forbid")
    source:str
    target:str
    kind:str=Field(min_length=1,max_length=64)
class Metric(BaseModel):
    model_config=ConfigDict(extra="forbid")
    key:str
    timestamp:datetime
    metric:str=Field(min_length=1,max_length=128)
    value:float
    unit:str=Field(min_length=1,max_length=64)
    @field_validator("timestamp")
    @classmethod
    def tz(cls,v):
        if v.tzinfo is None:raise ValueError("timezone required")
        return v
    @field_validator("value")
    @classmethod
    def finite(cls,v):
        if not math.isfinite(v):raise ValueError("finite numeric value required")
        return v
class Extraction(BaseModel):
    model_config=ConfigDict(extra="forbid")
    text:str
    tables:list[dict[str,Any]]=[]
    entities:list[Entity]=[]
    relations:list[Relation]=[]
    metrics:list[Metric]=[]

class Parser:
    def __init__(self,settings=None):
        from .config import Settings
        self.settings=settings or Settings()
    def parse(self,path,suffix):
        try:
            if suffix in (".txt",".csv"):
                with path.open(encoding="utf-8-sig") as stream:text=stream.read(self.settings.max_text_chars+1)
                if len(text)>self.settings.max_text_chars:raise IntegrationError("PARSER_LIMIT_EXCEEDED","Text parser size limit exceeded")
                tables=[]
                if suffix==".csv":
                    rows=[];cells=0
                    for row in csv.reader(io.StringIO(text)):
                        cells+=len(row)
                        if cells>self.settings.max_table_cells:raise IntegrationError("PARSER_LIMIT_EXCEEDED","Table cell limit exceeded")
                        rows.append(row)
                    if rows:tables=[{"name":"CSV","headers":rows[0],"rows":rows[1:]}]
                return Extraction(text=text,tables=tables)
            if suffix==".xlsx":
                from openpyxl import load_workbook
                book=load_workbook(path,read_only=True,data_only=True)
                tables=[]
                try:
                    for sheet in book:
                        if (sheet.max_row or 0)*(sheet.max_column or 0)>self.settings.max_table_cells:
                            raise IntegrationError("PARSER_LIMIT_EXCEEDED","Workbook dimension limit exceeded")
                        rows=[[v.isoformat() if isinstance(v,datetime) else str(v) if v is not None else "" for v in row] for row in sheet.values]
                        if rows:tables.append({"name":sheet.title,"headers":rows[0],"rows":rows[1:]})
                finally:book.close()
                text="\n\n".join(t["name"]+"\n"+" | ".join(t["headers"])+"\n"+"\n".join(" | ".join(r) for r in t["rows"]) for t in tables)
                return Extraction(text=text,tables=tables)
            # Reuse existing PDF/DOCX MarkItDown/cleaning code; disable implicit LLM/OCR defaults.
            from src.converter import FileConverter
            class DeterministicConverter(FileConverter):
                def _build_default_llm_client(self,llm_model=None):return None
            result=DeterministicConverter().convert_file(str(path),str(path.with_suffix(".parsed.md")))
            if result.get("status")!="success":raise IntegrationError("PARSER_REJECTION","Existing document parser rejected source")
            return Extraction(text=result.get("content",""))
        except IntegrationError:raise
        except (ImportError,ModuleNotFoundError):
            raise IntegrationError("UNSUPPORTED_DOCUMENT_FORMAT","Parser dependency unavailable") from None
        except Exception:raise IntegrationError("PARSER_REJECTION","Document parser rejected source") from None

def normalize(extraction,job,rules):
    data=extraction.model_dump()
    data["text"]="\n".join(line.rstrip() for line in extraction.text.replace("\r\n","\n").split("\n")).strip()
    metrics=list(data["metrics"])
    recognized_tables=set()
    for table in extraction.tables:
        headers=[str(x).strip().lower() for x in table["headers"]]
        for rule in rules:
            # Rules are configuration, not LLM decisions. Require exact schema and explicit unit.
            if rule["timestamp"].lower() not in headers or rule["value"].lower() not in headers:continue
            recognized_tables.add(table["name"])
            for i,row in enumerate(table["rows"]):
                if len(row)!=len(headers):raise IntegrationError("SCHEMA_VALIDATION_FAILED","Recognized metric row shape invalid")
                try:
                    metrics.append(Metric(key=stable_id(table["name"],rule["metric"],i),
                        timestamp=datetime.fromisoformat(str(row[headers.index(rule["timestamp"].lower())]).replace("Z","+00:00")),
                        metric=rule["metric"],value=float(row[headers.index(rule["value"].lower())]),unit=rule["unit"]).model_dump())
                except Exception:raise IntegrationError("SCHEMA_VALIDATION_FAILED","Recognized numeric row invalid") from None
    data["metrics"]=metrics
    if recognized_tables:
        # Raw recognized numeric series stay in typed artifacts/Timescale, never Qdrant.
        data["text"]="\n\n".join(t["name"]+"\n"+" | ".join(str(x) for x in t["headers"])+"\n"+"\n".join(" | ".join(str(x) for x in r) for r in t["rows"]) for t in extraction.tables if t["name"] not in recognized_tables)
    try:
        result=Extraction.model_validate(data)
        keys={x.key for x in result.entities}
        if len(keys)!=len(result.entities) or any(r.source not in keys or r.target not in keys for r in result.relations):
            raise ValueError("invalid entity references")
        return result
    except Exception:raise IntegrationError("SCHEMA_VALIDATION_FAILED","Typed extraction schema invalid") from None

class Pipeline:
    def run_report(self,job,path):
        from .report_adapter import REPORT_PARSER_VERSION,REPORT_EXTRACTION_VERSION
        self.repo.phase(job,"Parsing")
        pkg=self.reports.parse(path,job)
        checkpoint=self.repo.checkpoints(job);references={}
        for phase,write in (("TimescaleDB",self.reports.write_numeric),("Neo4j",self.reports.write_graph),("Qdrant",self.reports.write_vectors)):
            self.repo.phase(job,phase)
            if phase in checkpoint:references[phase]=checkpoint[phase];continue
            try:result=write(pkg)
            except IntegrationError:raise
            except Exception:raise IntegrationError("STORAGE_UNAVAILABLE",phase+" writer unavailable",True) from None
            self.repo.checkpoint(job,phase,result)
            references[phase]=result
        self.repo.finish(job,references,REPORT_PARSER_VERSION,REPORT_EXTRACTION_VERSION)

    def __init__(self,settings,repository,client,writers,parser=None,extractor=None,numeric_rules=(),reports=None):
        self.settings,self.repo,self.client,self.writers=settings,repository,client,writers
        self.reports=reports # ReportAdapter: workbooks with a registered report template skip the generic extraction
        self.parser=parser or Parser(settings)
        self.extractor=extractor # controlled service returns typed content; never receives writer tools
        self.numeric_rules=numeric_rules
    def run(self,job):
        s=self.settings
        root=Path(s.storage_root)/str(job["job_id"])
        # Original filename is never a directory or storage identity.
        path=root/("source-"+str(job["lease_token"])+Path(job["file_name"]).suffix.lower())
        self.repo.phase(job,"Downloading")
        digest=self.client.download(job,path)
        self.repo.set_hash(job,digest)
        self.repo.phase(job,"Validating")
        suffix=validate_file(path,job,s)
        if self.reports is not None and self.reports.matches(job):return self.run_report(job,path)
        checkpoint=self.repo.checkpoints(job)
        canonical=root/("extracted-"+str(job["lease_token"])+".json")
        if "Extraction" in checkpoint:
            try:
                canonical=Path(checkpoint["Extraction"]["artifact"])
                if canonical.resolve().parent!=root.resolve():raise ValueError()
                raw=canonical.read_bytes()
                if hashlib.sha256(raw).hexdigest()!=checkpoint["Extraction"]["sha256"]:raise ValueError()
                extraction=Extraction.model_validate_json(raw)
            except Exception:raise IntegrationError("CHECKPOINT_ARTIFACT_INVALID","Canonical extraction artifact unavailable") from None
        else:
            self.repo.phase(job,"Parsing")
            extraction=self.parser.parse(path,suffix)
            if self.extractor:
                self.repo.phase(job,"Extracting")
                try:extraction=Extraction.model_validate(self.extractor(extraction.model_dump()))
                except Exception:raise IntegrationError("SCHEMA_VALIDATION_FAILED","Semantic extraction schema invalid") from None
            self.repo.phase(job,"Normalizing")
            extraction=normalize(extraction,job,self.numeric_rules)
            raw=extraction.model_dump_json().encode()
            temp=root/("extracted-"+str(job["lease_token"])+".tmp")
            temp.write_bytes(raw)
            temp.replace(canonical)
            self.repo.checkpoint(job,"Extraction",{"sha256":hashlib.sha256(raw).hexdigest(),"artifact":str(canonical)})
        if not extraction.text and not extraction.metrics and not extraction.entities:
            raise IntegrationError("EMPTY_EXTRACTION","No validated knowledge extracted")
        meta=lineage(job)
        checkpoint=self.repo.checkpoints(job)
        references={}
        for phase,writer,items in (
            ("Storage",self.writers["storage"],extraction),
            ("Qdrant",self.writers.get("qdrant"),extraction.text),
            ("Neo4j",self.writers.get("neo4j"),(extraction.entities,extraction.relations)),
            ("TimescaleDB",self.writers.get("timescale"),extraction.metrics)):
            applicable=bool(items) if phase not in ("Storage","Neo4j") else phase=="Storage" or bool(items[0] or items[1])
            if not applicable:continue
            if writer is None:raise IntegrationError("STORAGE_UNAVAILABLE","Required "+phase+" writer unavailable",True)
            self.repo.phase(job,phase)
            if phase in checkpoint:references[phase]=checkpoint[phase];continue
            try:result=writer.write(items,meta,job)
            except IntegrationError:raise
            except Exception:raise IntegrationError("STORAGE_UNAVAILABLE",phase+" writer unavailable",True) from None
            self.repo.checkpoint(job,phase,result)
            references[phase]=result
        self.repo.finish(job,references,PARSER_VERSION,EXTRACTION_VERSION)
