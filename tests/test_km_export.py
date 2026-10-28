import io
import json
import zipfile
from pathlib import Path
from xml.dom import minidom

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pydantic import ValidationError

from src.km_export import builders, planner, service
from src.km_export.spec import FORMATS, ExportSpec
from src.km_export.store import ExportSettings, ExportStore, RateLimiter

SPEC = {
    "format": "pptx", "title": "Wi-Fi 8 重點整理", "subtitle": "依知識庫回答整理",
    "sections": [
        {"heading": "概觀", "bullets": [f"重點 {i}" for i in range(1, 9)], "paragraphs": ["802.11bn 目標是提升可靠度。"], "sources": ["WFA 規格 v1.0"]},
        {"heading": "比較", "table": {"columns": ["項目", "Wi-Fi 7", "Wi-Fi 8"], "rows": [[f"列{i}", "a|b", "=SUM(A1)"] for i in range(12)]}, "sources": ["IEEE 802.11bn D1.0"]},
        {"heading": "", "paragraphs": ["空標題的段落"]},
    ],
}


def spec(fmt: str) -> ExportSpec:
    return ExportSpec.model_validate({**SPEC, "format": fmt})


@pytest.mark.parametrize("fmt", FORMATS)
def test_every_format_builds_with_the_disclaimer_and_a_matching_name(fmt):
    data, media_type, name = builders.build(spec(fmt))
    assert data and name == f"Wi-Fi 8 重點整理.{fmt}" and media_type == builders.MEDIA_TYPES[fmt]
    if fmt in ("md", "csv", "html"):
        assert "請對照來源" in data.decode("utf-8")


def test_pptx_has_cover_split_bullets_and_split_table():
    from pptx import Presentation
    deck = Presentation(io.BytesIO(builders.build(spec("pptx"))[0]))
    assert len(deck.slides) == 1 + 2 + 2 + 1  # cover, 9 points over two slides, 12 rows over two slides, the untitled paragraph
    assert any(shape.has_table for slide in deck.slides for shape in slide.shapes)


def test_docx_is_valid_zip_with_well_formed_xml_and_escaped_text():
    s = ExportSpec.model_validate({**SPEC, "format": "docx", "title": "A & B <x>"})
    data = builders.build(s)[0]
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for part in z.namelist():
            minidom.parseString(z.read(part))  # every part is well-formed XML
        assert "A &amp; B &lt;x&gt;" in z.read("word/document.xml").decode("utf-8")


def test_xlsx_sheet_per_table_and_formula_text_stays_text():
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(builders.build(spec("xlsx"))[0]))
    assert wb.sheetnames == ["摘要", "比較"]
    assert str(wb["比較"]["C2"].value).startswith("'=")


def test_csv_has_bom_and_table_rows():
    text = builders.build(spec("csv"))[0].decode("utf-8")
    assert text.startswith("﻿") and "項目,Wi-Fi 7,Wi-Fi 8" in text


def test_markdown_escapes_pipes_and_html_escapes_markup():
    assert "a\\|b" in builders.build(spec("md"))[0].decode()
    s = ExportSpec.model_validate({**SPEC, "format": "html", "title": "<script>alert(1)</script>"})
    assert "<script>" not in builders.build(s)[0].decode()


def test_ragged_table_rows_are_squared():
    s = ExportSpec.model_validate({"format": "md", "title": "t", "sections": [{"table": {"columns": ["a", "b"], "rows": [["1"], ["1", "2", "3"]]}}]})
    assert "| 1 |  |" in builders.build(s)[0].decode()


def test_spec_rejects_unknown_format_and_empty_sections():
    with pytest.raises(ValidationError):
        ExportSpec.model_validate({**SPEC, "format": "exe"})
    with pytest.raises(ValidationError):
        ExportSpec.model_validate({"format": "md", "title": "t", "sections": []})


# ------------------------------------------------------------------ planner

def test_planner_reads_fenced_json_and_forces_the_requested_format():
    reply = "好的\n```json\n" + json.dumps({**SPEC, "format": "docx"}) + "\n```"
    got = planner.plan(lambda m: reply, question="q", answer="a", sources=["s"], fmt="pptx")
    assert got.format == "pptx" and got.sections[0].heading == "概觀"


def test_planner_retries_once_then_gives_a_plan_error():
    calls = []
    def bad(messages):
        calls.append(messages)
        return "not json"
    with pytest.raises(planner.PlanError, match="invalid_plan"):
        planner.plan(bad, question="", answer="a", sources=[])
    assert len(calls) == 2 and "不是有效" in calls[1][-1]["content"]


def test_planner_reports_an_unavailable_model():
    def down(messages):
        raise RuntimeError("503")
    with pytest.raises(planner.PlanError, match="llm_unavailable"):
        planner.plan(down, question="", answer="a", sources=[])


def test_planner_prompt_forbids_new_facts_and_carries_the_user_request():
    messages = planner.build_messages("Q?", "答案內容", ["文件 A"], "做成 5 頁簡報給主管", "pptx")
    assert "不可新增" in messages[0]["content"]
    assert "做成 5 頁簡報給主管" in messages[1]["content"] and "[1] 文件 A" in messages[1]["content"]


# ------------------------------------------------------------------ store

def make_store(tmp_path, clock):
    return ExportStore(ExportSettings(directory=tmp_path, ttl_seconds=100, max_files_per_owner=2), clock=clock)


def test_store_is_owner_only_and_expires(tmp_path):
    now = [1000.0]
    store = make_store(tmp_path, lambda: now[0])
    saved = store.save("a", b"data", "text/plain", "x.txt")
    assert store.open("a", saved["file_id"]) and store.open("b", saved["file_id"]) is None
    assert store.open("a", "../etc/passwd") is None
    now[0] += 101
    assert store.open("a", saved["file_id"]) is None
    store.purge()
    assert not list(tmp_path.iterdir())


def test_store_keeps_only_the_newest_files_per_owner(tmp_path):
    now = [1000.0]
    store = make_store(tmp_path, lambda: now[0])
    ids = []
    for _ in range(3):
        now[0] += 1
        ids.append(store.save("a", b"d", "text/plain", "x.txt")["file_id"])
    assert store.open("a", ids[0]) is None and store.open("a", ids[1]) and store.open("a", ids[2])


def test_rate_limiter_slides():
    now = [0.0]
    limiter = RateLimiter(lambda: now[0])
    assert [limiter.allow("u", 2) for _ in range(3)] == [True, True, False]
    now[0] += 3601
    assert limiter.allow("u", 2)


# ------------------------------------------------------------------ routes

@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("KM_EXPORT_DIR", str(tmp_path))
    service._store.cache_clear()
    service._limiter.__init__()
    state = {"allowed": True, "user": "u1"}
    app = FastAPI()
    app.include_router(service.router)

    @app.middleware("http")
    async def login(request, call_next):
        if request.headers.get("x-test-anon") != "1":
            request.state.km_identity = {"sub": state["user"], "instance_id": "test", "role": "user"}
        return await call_next(request)

    monkeypatch.setattr(service, "_may_export", lambda request: state["allowed"])
    monkeypatch.setattr(service, "_llm_complete", lambda: (lambda messages: json.dumps(SPEC)))
    test_client = TestClient(app)
    test_client.state = state
    return test_client


POST = {"x-km-export": "1"}


def test_capability_id_matches_the_catalog_and_is_shared_with_attachments():
    catalog = json.loads((Path(__file__).parents[1] / "config" / "km_capabilities.json").read_text("utf-8"))["capabilities"]
    entry = next(c for c in catalog if c["id"] == service.CAPABILITY)
    assert "匯出" in entry["label"] and "上傳" in entry["label"]
    from src.wifi8 import userfiles
    assert userfiles.CAPABILITY == service.CAPABILITY


def test_anonymous_is_refused_and_unauthorised_sees_it_switched_off(client):
    assert client.get("/api/export/config", headers={"x-test-anon": "1"}).status_code == 401
    client.state["allowed"] = False
    assert client.get("/api/export/config").json() == {"enabled": False}
    assert client.post("/api/export/plan", json={"answer": "a"}, headers=POST).status_code == 403
    assert client.post("/api/export/build", json={"spec": SPEC}, headers=POST).status_code == 403
    assert client.get("/api/export/files/" + "0" * 32).status_code == 403


def test_plan_build_download_round_trip(client):
    assert client.get("/api/export/config").json()["formats"] == list(FORMATS)
    planned = client.post("/api/export/plan", json={"question": "q", "answer": "a", "sources": ["s"], "format": "docx"}, headers=POST).json()
    assert planned["outline"] == ["概觀", "比較", "（無標題）"] and planned["spec"]["format"] == "docx"
    built = client.post("/api/export/build", json={"spec": planned["spec"]}, headers=POST)
    assert built.status_code == 201
    got = client.get("/api/export/files/" + built.json()["file_id"])
    assert got.status_code == 200 and zipfile.is_zipfile(io.BytesIO(got.content))
    assert "attachment" in got.headers["content-disposition"] and got.headers["content-security-policy"] == "sandbox"
    client.state["user"] = "someone-else"
    assert client.get("/api/export/files/" + built.json()["file_id"]).status_code == 404


def test_writes_need_the_same_origin_header(client):
    assert client.post("/api/export/plan", json={"answer": "a"}).status_code == 403
    assert client.post("/api/export/build", json={"spec": SPEC}).status_code == 403


def test_plan_errors_are_clear_and_the_hourly_limit_applies(client, monkeypatch):
    def down(messages):
        raise RuntimeError("503")
    monkeypatch.setattr(service, "_llm_complete", lambda: down)
    r = client.post("/api/export/plan", json={"answer": "a"}, headers=POST)
    assert r.status_code == 503 and "語言模型" in r.json()["detail"]
    monkeypatch.setenv("KM_EXPORT_PLANS_PER_HOUR", "1")
    assert client.post("/api/export/plan", json={"answer": "a"}, headers=POST).status_code == 429
    assert client.post("/api/export/plan", json={"answer": "a", "format": "exe"}, headers=POST).status_code in (422, 429)


def test_oversized_or_empty_answers_are_rejected(client):
    assert client.post("/api/export/plan", json={"answer": ""}, headers=POST).status_code == 422
    assert client.post("/api/export/plan", json={"answer": "x" * 20001}, headers=POST).status_code == 422


# ------------------------------------------------------------------ several answers in one report

def item(question, answer, sources=()):
    return service.PlanItem(question=question, answer=answer, sources=list(sources))


def test_merge_gives_each_distinct_source_one_report_wide_number_and_rewrites_the_labels():
    first = item("Q1", "甲 [S1] 乙 [S2]", ["[S1] 文件 A v1", "[S2] 文件 B v2"])
    second = item("Q2", "丙 [S1・我的檔案] 丁 [S3]", ["[S1] 文件 B v2", "[S2] 我的檔案.pdf"])  # [S1] here is document B; [S3] has no source
    pairs, sources = planner.merge_answers([first, second])
    assert sources == ["文件 A v1", "文件 B v2", "我的檔案.pdf"]
    assert pairs[0] == ("Q1", "甲 [1] 乙 [2]")
    assert pairs[1] == ("Q2", "丙 [2] 丁 [Q2-S3]")  # B is [2] in both answers; the unlisted label is marked, not guessed


def test_merge_keeps_answers_without_any_source_list():
    pairs, sources = planner.merge_answers([item("Q", "只有文字")])
    assert pairs == [("Q", "只有文字")] and sources == []


def test_multi_prompt_has_every_question_the_global_sources_and_the_conflict_rule():
    pairs, sources = planner.merge_answers([item("第一問", "答一 [S1]", ["[S1] 文件 A"]), item("第二問", "答二 [S1]", ["[S1] 文件 C"])])
    messages = planner.build_messages_multi(pairs, sources, "整合成一份報告", "docx")
    system, user = messages[0]["content"], messages[1]["content"]
    assert "不一致" in system and "不可自行挑一個" in system and "[Q1-S2]" in system
    assert "第一問" in user and "第二問" in user and "[1] 文件 A" in user and "[2] 文件 C" in user and "答二 [2]" in user


def test_planner_uses_the_multi_prompt_when_given_items():
    seen = []
    def complete(messages):
        seen.append(messages)
        return json.dumps(SPEC)
    planner.plan(complete, items=[item("Q1", "A1"), item("Q2", "A2")], fmt="pptx")
    assert "--- 第 2 則 ---" in seen[0][1]["content"]


def test_route_takes_up_to_five_answers_and_sends_the_merged_text(client, monkeypatch):
    seen = []
    monkeypatch.setattr(service, "_llm_complete", lambda: (lambda messages: (seen.append(messages), json.dumps(SPEC))[1]))
    items = [{"question": f"Q{i}", "answer": f"A{i} [S1]", "sources": [f"[S1] 文件 {i}"]} for i in range(1, 6)]
    r = client.post("/api/export/plan", json={"items": items, "format": "docx"}, headers=POST)
    assert r.status_code == 200 and r.json()["spec"]["format"] == "docx"
    assert "A5 [5]" in seen[0][1]["content"] and "[5] 文件 5" in seen[0][1]["content"]


def test_route_refuses_a_sixth_answer_with_a_clear_message(client):
    items = [{"answer": f"A{i}"} for i in range(6)]
    r = client.post("/api/export/plan", json={"items": items}, headers=POST)
    assert r.status_code == 422 and "最多匯出 5 則" in r.json()["detail"]


def test_route_refuses_too_much_text_across_the_selected_answers(client):
    items = [{"answer": "x" * 15000} for _ in range(3)]  # each is allowed alone; together they are over the cap
    r = client.post("/api/export/plan", json={"items": items}, headers=POST)
    assert r.status_code == 413 and "太長" in r.json()["detail"]


def test_a_request_with_neither_answer_nor_items_is_refused(client):
    assert client.post("/api/export/plan", json={"question": "q"}, headers=POST).status_code == 422


def test_several_answer_reports_have_their_own_smaller_hourly_limit(client, monkeypatch):
    monkeypatch.setenv("KM_EXPORT_MULTI_PLANS_PER_HOUR", "1")
    body = {"items": [{"answer": "a"}, {"answer": "b"}]}
    assert client.post("/api/export/plan", json=body, headers=POST).status_code == 200
    assert client.post("/api/export/plan", json=body, headers=POST).status_code == 429
    assert client.post("/api/export/plan", json={"answer": "a"}, headers=POST).status_code == 200  # a single answer is not affected


def test_config_tells_the_page_the_answer_limit(client):
    assert client.get("/api/export/config").json()["max_answers"] == 5
