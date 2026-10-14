import io,json,zipfile
from pathlib import Path
from dataclasses import replace
import httpx,pytest
from src.csit_integration_v1.models import IngestionRequest
from src.csit_integration_v1.client import CsitIntegrationClient
from src.csit_integration_v1.pipeline import Parser,Pipeline,Extraction,validate_file,normalize,lineage
from src.csit_integration_v1.writers import StorageWriter
from src.csit_integration_v1.errors import IntegrationError
from src.csit_integration_v1.worker import Worker

class FakeAuth:
    def headers(self):return {"Authorization":"Bearer fake-outbound-service"}
class FakeCsit:
    def __init__(self,payload,content=b"report text!"):
        self.payload=payload;self.content=content;self.download_status=200;self.callback_status=200
        self.downloads=0;self.callbacks=[];self.expected_job=None;self.timeout=False
    def __call__(self,r):
        assert r.headers.get("authorization")=="Bearer fake-outbound-service"
        if r.method=="GET":
            self.downloads+=1
            if self.timeout:raise httpx.ReadTimeout("synthetic timeout")
            return httpx.Response(self.download_status,content=self.content,headers={
                "Content-Type":self.payload["document"]["contentType"],
                "Content-Disposition":'attachment; filename="'+self.payload["document"]["fileName"]+'"'})
        p=json.loads(r.content)
        expected={'Processing':{'updatedAt','startedAt'},'Completed':{'updatedAt','completedAt'},'Failed':{'updatedAt'}}
        assert set(p['timestamps'])==expected[p['status']]
        assert not any(key in p for key in ['updatedAt','startedAt','completedAt','failedAt'])
        if p['status']=='Processing':
            from datetime import datetime
            assert datetime.fromisoformat(p['timestamps']['startedAt']).tzinfo is not None
        if p["jobId"]!=self.expected_job or p["sourceSystem"]!="CSIT" or p["documentId"]!=self.payload["documentId"] or p["version"]!=self.payload["version"]:
            return httpx.Response(409)
        self.callbacks.append(p)
        return httpx.Response(self.callback_status)

class RecordingWriter:
    def __init__(self):self.rows={};self.calls=0;self.fail_once=False
    def write(self,items,meta,job):
        self.calls+=1
        if self.fail_once:
            self.fail_once=False
            raise IntegrationError("STORAGE_UNAVAILABLE","Synthetic downstream outage",True)
        self.rows[(meta["documentId"],meta["version"])]=items
        return {"sourceVersion":str(job["job_id"])}

def prepare(repo,settings,payload,content=b"report text!",extractor=None):
    payload["document"]["sizeBytes"]=len(content)
    server=FakeCsit(payload,content)
    client=CsitIntegrationClient(settings,FakeAuth(),httpx.MockTransport(server))
    r=repo.submit(IngestionRequest(**payload));server.expected_job=str(r["job_id"])
    q,n,t=RecordingWriter(),RecordingWriter(),RecordingWriter()
    writers={"storage":StorageWriter(settings.storage_root),"qdrant":q,"neo4j":n,"timescale":t}
    rules=[{"timestamp":"timestamp","value":"throughput_mbps","metric":"throughput","unit":"Mbit/s"}]
    pipeline=Pipeline(settings,repo,client,writers,extractor=extractor,numeric_rules=rules)
    return Worker(settings,repo,client,pipeline),server,r,(q,n,t)

def test_stream_text_completed_lineage(repo,settings,payload):
    worker,server,r,(q,n,t)=prepare(repo,settings,payload)
    assert worker.process_one()
    final=repo.get(r["job_id"])
    assert final["status"]=="Completed" and final["parser_version"] and final["extraction_version"]
    assert q.calls==1 and n.calls==0 and t.calls==0
    normalized=json.loads((Path(settings.storage_root)/str(r["job_id"])/"normalized.json").read_text())
    assert normalized["lineage"]==lineage(r)
    assert set(final["storage_references"])=={"Storage","Qdrant"}
@pytest.mark.parametrize("status,retryable",[(404,False),(401,False),(403,False),(429,True),(500,True),(503,True),(302,False)])
def test_download_failures(repo,settings,payload,status,retryable):
    worker,server,r,writers=prepare(repo,settings,payload)
    server.download_status=status;worker.process_one()
    final=repo.get(r["job_id"])
    assert final["status"]=="Failed"
    assert bool(final["next_retry_at"])==retryable
    assert all(w.calls==0 for w in writers)
def test_timeout_retry(repo,settings,payload):
    worker,server,r,_=prepare(repo,settings,payload)
    server.timeout=True;worker.process_one()
    assert repo.get(r["job_id"])["next_retry_at"]
    server.timeout=False;worker.process_one()
    assert repo.get(r["job_id"])["status"]=="Completed"
def test_corrupt_signature(repo,settings,payload):
    payload["document"]={"fileName":"report.pdf","contentType":"application/pdf","sizeBytes":12}
    worker,server,r,_=prepare(repo,settings,payload);worker.process_one()
    final=repo.get(r["job_id"])
    assert final["status"]=="Failed" and final["last_error_code"]=="INVALID_FILE_SIGNATURE" and final["next_retry_at"] is None
def test_unsupported_document(repo,settings,payload):
    payload["document"].update(fileName="report.exe",contentType="application/octet-stream")
    worker,server,r,_=prepare(repo,settings,payload);worker.process_one()
    assert repo.get(r["job_id"])["last_error_code"]=="UNSUPPORTED_DOCUMENT_FORMAT"
@pytest.mark.parametrize("kind",["mime","filename","size","max"])
def test_response_validation(repo,settings,payload,kind,tmp_path):
    job=repo.submit(IngestionRequest(**payload))
    def handler(r):
        return httpx.Response(200,content=b"report text!",headers={"Content-Type":"application/pdf" if kind=="mime" else "text/plain",
            "Content-Disposition":'attachment; filename="../report.txt"' if kind=="filename" else 'attachment; filename="report.txt"'})
    if kind=="size":job["size_bytes"]=11
    if kind=="max":settings=replace(settings,max_file_size=5)
    client=CsitIntegrationClient(settings,FakeAuth(),httpx.MockTransport(handler))
    with pytest.raises(IntegrationError):client.download(job,tmp_path/"source.txt")
    assert not (tmp_path/"source.part").exists()
def test_numeric_not_vector(repo,settings,payload):
    payload["document"].update(fileName="report.csv",contentType="text/csv")
    worker,server,r,(q,n,t)=prepare(repo,settings,payload,b"timestamp,throughput_mbps\n2026-10-07T08:30:00+08:00,125.5\n")
    worker.process_one()
    assert repo.get(r["job_id"])["status"]=="Completed" and t.calls==1 and q.calls==0
    metric=next(iter(t.rows.values()))[0]
    assert metric.metric=="throughput" and metric.value==125.5 and metric.unit=="Mbit/s"
def test_general_table_not_numeric(repo,settings,payload):
    payload["document"].update(fileName="report.csv",contentType="text/csv")
    worker,server,r,(q,n,t)=prepare(repo,settings,payload,b"label,value\nresult,42\n")
    worker.process_one()
    assert repo.get(r["job_id"])["status"]=="Completed" and t.calls==0 and q.calls==1
def test_typed_entities_route(repo,settings,payload):
    def extract(data):
        data["entities"]=[{"key":"a","kind":"Device","name":"DUT"},{"key":"b","kind":"Issue","name":"timeout"}]
        data["relations"]=[{"source":"a","target":"b","kind":"HAS_ISSUE"}]
        return data
    worker,server,r,(q,n,t)=prepare(repo,settings,payload,extractor=extract)
    worker.process_one()
    assert repo.get(r["job_id"])["status"]=="Completed" and n.calls==1 and t.calls==0
def test_invalid_semantic_schema(repo,settings,payload):
    worker,server,r,w=prepare(repo,settings,payload,extractor=lambda _:{"text":"x","arbitraryDb":"neo4j"})
    worker.process_one()
    assert repo.get(r["job_id"])["status"]=="Failed" and all(x.calls==0 for x in w)
def test_partial_failure_checkpoint_retry(repo,settings,payload):
    calls=[]
    def extract(data):calls.append(1);return data
    worker,server,r,(q,n,t)=prepare(repo,settings,payload,extractor=extract)
    q.fail_once=True;worker.process_one()
    assert repo.get(r["job_id"])["status"]=="Failed"
    worker.process_one()
    assert repo.get(r["job_id"])["status"]=="Completed" and q.calls==2 and len(calls)==1
    assert len(q.rows)==1
    assert "Storage" in repo.checkpoints(r)
def test_callbacks_exact_completed_retry(repo,settings,payload):
    worker,server,r,_=prepare(repo,settings,payload);worker.process_one()
    server.callback_status=503
    assert worker.callback_one()
    assert repo.get(r["job_id"])["status"]=="Completed"
    downloads=server.downloads
    server.callback_status=200
    while worker.callback_one():pass
    assert not worker.process_one() and server.downloads==downloads
    assert [x["status"] for x in server.callbacks]==["Processing","Processing","Completed"]
    p=server.callbacks[-1]
    assert set(p)=={"jobId","sourceSystem","documentId","version","status","timestamps"}
    assert set(p["timestamps"])=={"updatedAt","completedAt"}
def test_failed_callback(repo,settings,payload):
    worker,server,r,_=prepare(repo,settings,payload);server.download_status=404;worker.process_one()
    while worker.callback_one():pass
    p=server.callbacks[-1]
    assert p["status"]=="Failed" and set(p["error"])=={"code","message"} and "completedAt" not in p["timestamps"]
@pytest.mark.parametrize("field",["jobId","sourceSystem","documentId","version"])
def test_callback_correlation_wrong_job(repo,settings,payload,field):
    worker,server,r,_=prepare(repo,settings,payload);worker.process_one()
    original=worker.client.callback
    def wrong(p):
        bad=dict(p);bad[field]="wrong" if field!="version" else 100
        original(bad)
    worker.client.callback=wrong
    worker.callback_one()
    assert server.callbacks==[]
    with repo.connect() as c:
        rows=c.execute("SELECT status,last_error_code FROM km_csit_callback_deliveries ORDER BY sequence").fetchall()
    assert rows[0]["status"]=="Pending"
    for _ in range(2):worker.callback_one()
    with repo.connect() as c:
        failed=c.execute("SELECT status FROM km_csit_callback_deliveries ORDER BY sequence LIMIT 1").fetchone()
    assert failed["status"]=="Failed" and server.callbacks==[]
def test_callback_timeout_only_delivery(repo,settings,payload):
    worker,server,r,_=prepare(repo,settings,payload);worker.process_one()
    old=worker.client.callback
    def timeout(_):raise IntegrationError("CSIT_CALLBACK_NETWORK_ERROR","Synthetic timeout",True)
    worker.client.callback=timeout;worker.callback_one()
    assert repo.get(r["job_id"])["status"]=="Completed"
    worker.client.callback=old
    while worker.callback_one():pass
    assert server.downloads==1
def test_archive_bomb(settings,payload,tmp_path):
    payload["document"].update(fileName="report.xlsx",contentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    path=tmp_path/"source.xlsx"
    with zipfile.ZipFile(path,"w",zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml","x");z.writestr("xl/workbook.xml","x"*100000)
    with pytest.raises(IntegrationError,match="ARCHIVE_LIMIT"):validate_file(path,{"file_name":"report.xlsx","content_type":payload["document"]["contentType"]},settings)
def test_xlsx_table_discovery(settings,tmp_path):
    from openpyxl import Workbook
    p=tmp_path/"test.xlsx";w=Workbook()
    w.active.title="General";w.active.append(["label","value"]);w.active.append(["test",42])
    w.create_sheet("Throughput");w["Throughput"].append(["timestamp","throughput_mbps"]);w["Throughput"].append(["2026-10-07T08:30:00+08:00",80])
    w.save(p)
    extraction=Parser().parse(p,".xlsx")
    assert len(extraction.tables)==2
    result=normalize(extraction,{},[{"timestamp":"timestamp","value":"throughput_mbps","metric":"throughput","unit":"Mbit/s"}])
    assert len(result.metrics)==1 and "42" in result.text and "80" not in result.text
def test_docx_reuse_parser(tmp_path):
    from docx import Document
    p=tmp_path/"report.docx";d=Document();d.add_paragraph("Test Report passed");d.save(p)
    assert "Test Report" in Parser().parse(p,".docx").text
def test_pdf_reuse_parser(tmp_path):
    from reportlab.pdfgen.canvas import Canvas
    p=tmp_path/"report.pdf";c=Canvas(str(p));c.drawString(50,700,"Test Report passed");c.save()
    assert "Test Report" in Parser().parse(p,".pdf").text
def test_no_agent_routing():
    import inspect
    source=inspect.getsource(Pipeline.run)
    assert 'self.writers["storage"]' in source and "self.writers.get" in source
    assert "agent" not in source.lower()
def test_logs_no_token_or_file_content(repo,settings,payload,caplog):
    import logging
    caplog.set_level(logging.INFO)
    worker,server,r,_=prepare(repo,settings,payload);worker.process_one();worker.callback_one()
    assert "fake-outbound-service" not in caplog.text and "report text!" not in caplog.text
    assert settings.db_url not in caplog.text and "test-only-outbound" not in caplog.text

def test_sparse_xlsx_dimension_bomb(settings,tmp_path):
    from openpyxl import Workbook
    p=tmp_path/"sparse.xlsx";w=Workbook();w.active.cell(row=1000000,column=10000,value="x");w.save(p)
    with pytest.raises(IntegrationError,match="PARSER_LIMIT_EXCEEDED"):Parser(settings).parse(p,".xlsx")

def test_csv_cell_limit(settings,tmp_path):
    p=tmp_path/"large.csv";p.write_text("a,b\n"+"1,2\n"*6)
    with pytest.raises(IntegrationError,match="PARSER_LIMIT_EXCEEDED"):Parser(replace(settings,max_table_cells=10)).parse(p,".csv")
