#!/usr/bin/env python3
"""Generate the WP2 CSIT contract from Patty's REST API workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

PRIORITY_CATEGORIES = {
    "Module",
    "Record / Project",
    "Case",
    "Task",
    "Task File",
    "Task Report",
    "Case File",
    "File",
}
KNOWN_GAPS = [
    "CSIT_OPENAPI_RETRIEVAL_PENDING",
    "CSIT_SERVICE_AUTH_CONTRACT_MISSING",
    "CSIT_BOOKING_CONTRACT_PENDING",
    "CSIT_VALIDATION_REQUEST_CONTRACT_PENDING",
]


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def generate(source: Path) -> dict[str, object]:
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook["API清單"] if "API清單" in workbook.sheetnames else workbook.active
    header_row = next(sheet.iter_rows(values_only=True))
    headers = {_text(value): index for index, value in enumerate(header_row)}
    required = {"分類", "Method", "Endpoint"}
    if required.issubset(headers):
        start_row = 2
        field = lambda values, name: values[headers[name]] if name in headers and headers[name] < len(values) else None
        description_name = "用途" if "用途" in headers else "說明"
        auth_name = "Auth" if "Auth" in headers else "Authorization"
        parameter_name = "Query / Path 參數" if "Query / Path 參數" in headers else "Query / Path / Body 參數"
        response_name = "Response / 回傳"
        controller_name = "Controller"
    else:
        # Legacy inventory has a title and a four-row preamble.
        header_row = next(sheet.iter_rows(min_row=4, max_row=4, values_only=True))
        headers = {_text(value): index for index, value in enumerate(header_row)}
        start_row = 5
        field = lambda values, name: values[headers[name]] if name in headers and headers[name] < len(values) else None
        description_name = "說明"
        auth_name = "Authorization"
        parameter_name = "Query / Path / Body 參數"
        response_name = "Response / 回傳"
        controller_name = "Controller"
    rows = []
    for values in sheet.iter_rows(min_row=start_row, values_only=True):
        category = _text(field(values, "分類"))
        method = _text(field(values, "Method")).upper()
        endpoint = _text(field(values, "Endpoint"))
        controller = _text(field(values, controller_name))
        parameters = _text(field(values, parameter_name))
        response = _text(field(values, response_name))
        auth = _text(field(values, auth_name))
        description = _text(field(values, description_name))
        if category not in PRIORITY_CATEGORIES:
            continue
        if not endpoint or not method:
            continue
        rows.append(
            {
                "category": category,
                "controller": controller,
                "method": method,
                "path": endpoint,
                "parameters": parameters,
                "response": response,
                "authorization": auth,
                "description": description,
                "integration_stage": "WP2-A/B" if method == "GET" else "WP2-D-DEFERRED",
                "read_only_allowed": method == "GET",
            }
        )
    if not rows:
        raise ValueError("no WP2 priority rows found in workbook")
    return {
        "schema": "wp2.csit-rest-contract.v1",
        "source_file": source.name,
        "source_sha256": digest,
        "scope": ["Module", "Record / Project", "Case", "Task", "File / Test Report"],
        "read_only_boundary": "WP2-B permits only mapped GET endpoints",
        "known_gaps": KNOWN_GAPS,
        "endpoints": rows,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    result = generate(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"result": "PASS", "endpoint_count": len(result["endpoints"]), "source_sha256": result["source_sha256"]}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
